"""Read a supported file and return its text, with every refusal spelled out.

Supported: `.txt .md` (UTF-8), `.json` (pretty-printed), `.yaml`/`.yml` (`safe_load`, then
text), `.docx` (python-docx paragraphs and tables) and `.pdf` (pypdf page text). The file is
only read, never run. A symlink, a non-regular file, one over 10 MB, an unknown extension and
an unreadable or encrypted document are refused with a clear `StoreValidationError`. The
extracted text is capped at 2 MB so a huge document can't blow up the chunker or the index.
"""

import io
import json
import os
import stat
from dataclasses import dataclass
from pathlib import Path

from ..errors import StoreValidationError

MAX_FILE_BYTES = 10 * 1024 * 1024       # 10 MB on disk
MAX_TEXT_CHARS = 2 * 1024 * 1024        # 2 MB of extracted text
MAX_NAME = 120

_EXT_KIND = {
    ".txt": "txt", ".md": "md", ".json": "json",
    ".yaml": "yaml", ".yml": "yaml", ".docx": "docx", ".pdf": "pdf",
}
SUPPORTED = ".txt .md .json .yaml .yml .docx .pdf"


@dataclass(frozen=True)
class Extracted:
    """The outcome of reading one file: its display name, kind, byte size and text."""
    name: str
    kind: str
    bytes: int
    text: str


def extract(path_text: str) -> Extracted:
    """Read the file at `path_text` and return its `Extracted` text, or raise."""
    text = (path_text or "").strip()
    if not text:
        raise StoreValidationError("Enter the path to a file to add.")
    path = Path(text).expanduser()

    if path.is_symlink():
        raise StoreValidationError("That path is a symlink; add the real file instead.")
    if not path.exists():
        raise StoreValidationError("No file at that path.")
    if not path.is_file():
        raise StoreValidationError("That isn't a regular file.")
    kind = _EXT_KIND.get(path.suffix.lower())
    if kind is None:
        raise StoreValidationError(f"Unsupported file type. Use one of: {SUPPORTED}.")

    data = _read_regular(path)
    raw = _EXTRACTORS[kind](data)
    return Extracted(name=path.name[:MAX_NAME], kind=kind, bytes=len(data),
                     text=raw[:MAX_TEXT_CHARS])


def _read_regular(path: Path) -> bytes:
    """The file's bytes, read without following a symlink and refused over the size cap."""
    flags = os.O_RDONLY | getattr(os, "O_NOFOLLOW", 0) | getattr(os, "O_BINARY", 0)
    try:
        fd = os.open(path, flags)
    except OSError:
        # ELOOP (a symlink slipped in) or a permission problem: both are "couldn't read".
        raise StoreValidationError("Couldn't read that file.") from None
    try:
        info = os.fstat(fd)
        if not stat.S_ISREG(info.st_mode):
            raise StoreValidationError("That isn't a regular file.")
        if info.st_size > MAX_FILE_BYTES:
            raise StoreValidationError("That file is too large (over 10 MB).")
        chunks = []
        total = 0
        while True:
            block = os.read(fd, 1 << 20)
            if not block:
                break
            total += len(block)
            if total > MAX_FILE_BYTES:
                raise StoreValidationError("That file is too large (over 10 MB).")
            chunks.append(block)
    finally:
        os.close(fd)
    return b"".join(chunks)


# --- per-kind extractors -------------------------------------------------------------------

def _from_text(data: bytes) -> str:
    return data.decode("utf-8", errors="replace")


def _from_json(data: bytes) -> str:
    try:
        obj = json.loads(data.decode("utf-8", errors="replace"))
    except (ValueError, UnicodeError):
        raise StoreValidationError("That file isn't valid JSON.") from None
    return json.dumps(obj, indent=2, ensure_ascii=False, sort_keys=False)


def _from_yaml(data: bytes) -> str:
    import yaml

    try:
        obj = yaml.safe_load(data.decode("utf-8", errors="replace"))
    except yaml.YAMLError:
        raise StoreValidationError("That file isn't valid YAML.") from None
    if obj is None:
        return ""
    if isinstance(obj, (dict, list)):
        return yaml.safe_dump(obj, default_flow_style=False, allow_unicode=True, sort_keys=False)
    return str(obj)


def _from_docx(data: bytes) -> str:
    from docx import Document

    try:
        document = Document(io.BytesIO(data))
    except Exception:
        raise StoreValidationError("Couldn't read that Word file.") from None
    parts = [p.text for p in document.paragraphs if p.text and p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            line = " | ".join(c for c in cells if c)
            if line:
                parts.append(line)
    return "\n\n".join(parts)


def _from_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception:
        raise StoreValidationError("Couldn't read that PDF.") from None
    if reader.is_encrypted:
        # A blank owner password may still open it; try, then give up cleanly.
        try:
            if reader.decrypt("") <= 0:  # 0 = failed
                raise StoreValidationError("That PDF is encrypted; add an unencrypted copy.")
        except StoreValidationError:
            raise
        except Exception:
            raise StoreValidationError("That PDF is encrypted; add an unencrypted copy.") from None
    parts = []
    try:
        for page in reader.pages:
            text = page.extract_text() or ""
            if text.strip():
                parts.append(text)
    except Exception:
        raise StoreValidationError("Couldn't read that PDF.") from None
    return "\n\n".join(parts)


_EXTRACTORS = {
    "txt": _from_text,
    "md": _from_text,
    "json": _from_json,
    "yaml": _from_yaml,
    "docx": _from_docx,
    "pdf": _from_pdf,
}
