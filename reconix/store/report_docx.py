"""Render `report.report_data()` as a Word (.docx) document.

Optional: enabled only when python-docx is installed (`HAS_DOCX`). Writes a binary file
straight to `path`. Evidence is already masked in the data.
"""

import importlib.util
from typing import Any, Dict

HAS_DOCX = importlib.util.find_spec("docx") is not None


def _meta_rows(data: Dict[str, Any]):
    return [
        ("Client", data.get("client") or "—"),
        ("Engagement ID", data.get("assessment_id") or "—"),
        ("Assessment type", data.get("mode") or "—"),
        ("Report date", (data.get("generated_at") or "").split("T")[0] or "—"),
        ("Report version", data.get("report_version") or "1.0"),
        ("Standards", " · ".join(data.get("methodology", [])) or "—"),
    ]


def write_docx(data: Dict[str, Any], path) -> None:
    """Write the full report as a .docx file at `path`."""
    from docx import Document

    findings = data["findings"]
    doc = Document()
    doc.add_heading("Security Assessment Report", level=0)
    doc.add_paragraph(f"AI-Assisted Authorized Penetration Test — {data.get('template', '')}")

    meta = doc.add_table(rows=0, cols=2)
    meta.style = "Light List Accent 1"
    for label, value in _meta_rows(data):
        cells = meta.add_row().cells
        cells[0].text, cells[1].text = label, str(value)
    doc.add_paragraph("CONFIDENTIAL — authorized recipients only.")

    doc.add_heading("Executive summary", level=1)
    doc.add_paragraph(data.get("summary", ""))

    doc.add_heading("Findings summary", level=1)
    table = doc.add_table(rows=1, cols=6)
    table.style = "Light Grid Accent 1"
    for i, head in enumerate(["ID", "Severity", "CVSS", "Status", "Category", "Finding"]):
        table.rows[0].cells[i].text = head
    for f in findings:
        cells = table.add_row().cells
        cells[0].text = f["id"]
        cells[1].text = f["severity"]
        cells[2].text = f"{f.get('cvss_score') or ''}"
        cells[3].text = f.get("status", "")
        cells[4].text = f.get("category", "")
        cells[5].text = f["title"]

    doc.add_heading("Detailed findings", level=1)
    for f in findings:
        doc.add_heading(f"{f['id']} · {f['title']}", level=2)
        maps = list(f.get("owasp", [])) + list(f.get("cwe", [])) + list(f.get("cve", []))
        doc.add_paragraph(
            f"Severity: {f['severity']} · CVSS {f.get('cvss_score') or '—'} · "
            f"{f.get('category', '')} · {f.get('validation', '')} · {f.get('status', 'open')}")
        doc.add_paragraph(f"Affected asset: {f.get('affected_url', '')}")
        doc.add_paragraph(f["description"])
        if f.get("evidence"):
            evidence = doc.add_paragraph("\n".join(f["evidence"]))
            evidence.style = "Intense Quote"
        doc.add_paragraph(f"Impact: {f['impact']}")
        doc.add_paragraph(f"Remediation: {f['remediation']}")
        doc.add_paragraph(f"Mapping: {', '.join(maps) or '—'}")

    doc.save(str(path))
