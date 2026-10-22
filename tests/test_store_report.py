"""Reports: PDF, JSON and DOCX files only, once the assessment is complete."""

import json
import re
import zipfile

import pytest

from reconix import store
from reconix.models.scope import PolicyVerdict
from reconix.store import lists, report
from reconix.store.report_docx import write_docx
from reconix.store.report_html import render_html

from .support import TEST_PASSWORD, add_finding, run_to

# What XML 1.0 can't hold (tab, LF and CR are fine)
CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")


def test_no_report_before_the_assessment_is_complete():
    run_to("approval:approval-001")
    with pytest.raises(store.StoreValidationError):
        store.generate_report("json")


def test_report_formats_and_their_availability():
    run_to(None)
    formats = {f.id: f.available for f in store.list_report_formats()}
    assert list(formats) == ["pdf", "json", "docx"]   # nothing else is saved (no HTML, CSV…)
    assert formats["json"]
    assert formats["pdf"] == report.HAS_PDF      # PDF only with WeasyPrint or a Chrome
    assert formats["docx"] == report.HAS_DOCX    # DOCX only when python-docx is installed
    for fid, available in formats.items():
        if not available:
            with pytest.raises(store.StoreValidationError):
                store.generate_report(fid)
    for gone in ("nope", "html", "markdown", "csv", "sarif"):
        with pytest.raises(store.StoreValidationError):
            store.generate_report(gone)


@pytest.mark.skipif(not report.HAS_PDF, reason="no PDF engine (WeasyPrint or Chrome) available")
def test_pdf_report_is_the_styled_report_as_a_pdf(reports_dir):
    run_to(None)
    shown = store.generate_report("pdf")
    path = reports_dir / "RCX-DEMO-001.pdf"
    assert shown == str(path)
    assert path.read_bytes()[:5] == b"%PDF-"


@pytest.mark.skipif(not report.HAS_DOCX, reason="python-docx not installed")
def test_docx_report_is_written_as_a_binary_file(reports_dir):
    run_to(None)
    shown = store.generate_report("docx")
    path = reports_dir / "RCX-DEMO-001.docx"
    assert shown == str(path)
    assert path.read_bytes()[:2] == b"PK"        # a .docx is a zip archive


def _docx_text(path) -> str:
    with zipfile.ZipFile(path) as docx:
        return docx.read("word/document.xml").decode("utf-8")


def test_json_report_is_written_and_recorded(reports_dir):
    run_to(None)
    shown = store.generate_report("json")
    path = reports_dir / "RCX-DEMO-001.json"
    assert shown == str(path)
    data = json.loads(path.read_text())
    assert data["assessment_id"] == "RCX-DEMO-001"
    assert [f["id"] for f in data["findings"]] == ["001", "002", "003"]
    assert data["blocked_requests"] == [
        {"request": "GET /admin", "reason": "Path outside approved scope"}]
    assert [a["risk"] for a in data["approvals"]] == ["MEDIUM", "HIGH"]
    assert store.get_run().report_path == shown
    assert store.phase_progress()["report"] == 100
    assert store.list_events()[-1].kind == "report.generated"


def _docx_parts(path):
    """The paragraphs (style, text) and the tables (rows of cell texts) of a .docx file."""
    import docx

    doc = docx.Document(str(path))
    paragraphs = [(p.style.name, p.text) for p in doc.paragraphs if p.text.strip()]
    tables = [[[cell.text for cell in row.cells] for row in table.rows] for table in doc.tables]
    return paragraphs, tables


@pytest.mark.skipif(not report.HAS_DOCX, reason="python-docx not installed")
def test_docx_report_has_every_section(reports_dir, monkeypatch):
    """The Word report follows the PDF: its sections, order, headings and content."""
    run_to(None)
    monkeypatch.setattr(lists.current().findings[2], "severity", "CRITICAL")   # 003 first
    store.triage_finding("002", note="Retested: the marker is encoded now.")
    store.generate_report("docx")
    paragraphs, tables = _docx_parts(reports_dir / "RCX-DEMO-001.docx")
    headings = [text for style, text in paragraphs if style.startswith("Heading")]
    assert headings[:5] == ["1 Executive summary", "2 Findings summary",
                            "3 Approved scope & limitations", "4 Detailed findings",
                            "003 · Session cookie missing SameSite"]     # worst first
    assert [h for h in headings if h[0].isdigit() and " · " in h] == [
        "003 · Session cookie missing SameSite", "001 · Broken object-level authorization",
        "002 · Reflected input in search page"]
    assert headings[-3:] == ["5 Policy & governance", "Blocked by policy · 1",
                             "Human approvals · 2"]
    html = render_html(store.report_data())                 # the PDF has the same headings
    for heading in ("Executive summary", "Findings summary", "Approved scope &amp; limitations",
                    "Detailed findings", "Policy &amp; governance"):
        assert f"<h2>{heading}</h2>" in html
    text = "\n".join(t for _s, t in paragraphs)
    assert "The highest severity is Critical, led by Session cookie missing SameSite" in text
    assert "Retested: the marker is encoded now." in text                  # the analyst's note
    assert "Testing was limited to the approved actions" in text            # the disclaimer
    assert "References: OWASP A01:2021" in text
    assert "CONFIDENTIAL — AUTHORIZED RECIPIENTS ONLY" in text
    assert [["Critical", "High", "Medium", "Low", "Total"],
            ["1", "0", "2", "0", "3"]] in tables                            # severity counts
    summary = next(t for t in tables if t[0] == ["ID", "Finding", "Severity", "CVSS", "Status",
                                                  "Category"])
    assert [row[0] for row in summary[1:]] == ["003", "001", "002"]
    scope = dict(next(t for t in tables if t[0][0] == "Approval status"))
    assert scope["Approval status"] == "Approved" and scope["Excluded paths"] == "/admin"
    assert scope["Standards"] == "OWASP WSTG · OWASP Top 10 (2021)"
    approvals = next(t for t in tables if t[0][0] == "Risk")
    assert [row[0] for row in approvals[1:]] == ["Medium", "High"]
    assert "•••" in _docx_text(reports_dir / "RCX-DEMO-001.docx")      # evidence stays masked


@pytest.mark.skipif(not report.HAS_DOCX, reason="python-docx not installed")
def test_control_characters_never_stop_the_docx_report(reports_dir, tmp_path):
    """An agent finding and a blocked request holding ESC / NUL: the report is still
    written, and none of those characters reaches it (or the JSON)."""
    run_to(None)
    add_finding("Banner \u001b[31mred\u001b[0m\u0000 leak", "/\u001b[2J",
                description="NUL \u0000 and BEL \u0007 here",
                evidence=["Server: \u001b[2Jnginx\u0000"])
    lists.current().verdicts.append(
        PolicyVerdict("GET", "/a\x00b", False, "Path \x1b[1moutside\x00 approved scope"))
    store.generate_report("docx")
    store.generate_report("json")
    xml = _docx_text(reports_dir / "RCX-DEMO-001.docx")
    assert "Banner [31mred[0m leak" in xml and "Path [1moutside approved scope" in xml
    assert not CONTROL.search(xml)
    assert not CONTROL.search((reports_dir / "RCX-DEMO-001.json").read_text("utf-8"))


@pytest.mark.skipif(not report.HAS_DOCX, reason="python-docx not installed")
def test_a_value_the_docx_cannot_hold_is_the_stores_verdict(tmp_path):
    run_to(None)
    data = store.report_data()
    data["findings"][0]["title"] = "bad \x1b title"            # past report_data's cleaning
    with pytest.raises(store.StoreValidationError):
        write_docx(data, tmp_path / "x.docx")


def test_report_data_carries_cvss_category_and_engagement_metadata():
    run_to(None)
    data = store.report_data()
    assert data["client"] == "Example"                   # derived from staging.example.com
    assert data["mode"] == "Grey-box, Authorized"
    assert data["methodology"] == ["OWASP WSTG", "OWASP Top 10 (2021)"]
    assert data["window"]["started"] and data["window"]["finished"]
    first = next(f for f in data["findings"] if f["id"] == "001")
    assert first["cvss_score"] == 6.5
    assert first["cvss_vector"].startswith("CVSS:3.1/")
    assert first["category"] == "Web App"
    assert first["cwe"] == ["CWE-639"]
    assert first["owasp"] == ["OWASP A01:2021", "OWASP API1:2023"]


def test_the_html_the_pdf_prints_is_a_full_styled_document():
    run_to(None)
    text = render_html(store.report_data())
    assert text.startswith("<!doctype html>")
    assert "Security Assessment Report" in text          # masthead title
    assert "RCX-DEMO-001" in text                         # engagement id
    assert "Broken object-level authorization" in text   # a finding
    assert "Executive summary" in text and "Detailed findings" in text
    assert "CVSS 6.5" in text                             # CVSS on the finding
    assert "OWASP WSTG" in text and "Example" in text     # methodology + client cover
    assert "•••" in text                                  # evidence stays masked


def test_the_html_the_pdf_prints_escapes_dynamic_text(monkeypatch):
    run_to(None)
    finding = store.list_findings()[0]
    monkeypatch.setattr(finding, "description", '<script>alert("x")</script>')
    text = render_html(store.report_data())
    assert "<script>alert" not in text                    # injected markup is escaped
    assert "&lt;script&gt;" in text


def test_reports_never_contain_the_test_account(reports_dir):
    run_to(None)
    store.generate_report("json")
    texts = [(reports_dir / "RCX-DEMO-001.json").read_text(), render_html(store.report_data())]
    if report.HAS_DOCX:
        store.generate_report("docx")
        texts.append(_docx_text(reports_dir / "RCX-DEMO-001.docx"))
    for text in texts:
        assert TEST_PASSWORD not in text
        assert "demo.tester" not in text


@pytest.mark.parametrize("request_text", [
    store.DEMO_REQUEST,                                   # password + one-time code
    "test the API at https://api.example.com/v1",         # a session cookie
])
def test_no_report_format_contains_the_login(reports_dir, request_text):
    run_to(None, request_text)
    formats = ["json"]
    if report.HAS_DOCX:
        formats.append("docx")
    for format_id in formats:
        store.generate_report(format_id)
    for path in reports_dir.iterdir():
        text = _docx_text(path) if path.suffix == ".docx" else path.read_text()
        assert TEST_PASSWORD not in text, path.name
        assert "demo.tester" not in text, path.name
    html = render_html(store.report_data())                   # what the PDF prints
    assert TEST_PASSWORD not in html and "demo.tester" not in html


def test_a_triage_note_is_masked_in_the_report():
    run_to(None)
    store.triage_finding("002", note=f"retested with password={TEST_PASSWORD}")
    note = next(f for f in store.report_data()["findings"] if f["id"] == "002")["triage_note"]
    assert note.startswith("retested with")
    assert TEST_PASSWORD not in note


def test_a_network_report_lists_the_approved_ports(reports_dir):
    run_to(None, "port scan 192.0.2.10")
    scope = store.report_data()["scope"]
    assert scope["kind"] == "network"
    assert scope["allowed_ports"] == [22, 80, 443, 8080, 8443]
    assert "<dt>Allowed ports</dt>" in render_html(store.report_data())
    if report.HAS_DOCX:
        store.generate_report("docx")
        text = _docx_text(reports_dir / "RCX-DEMO-001.docx")
        assert "Allowed ports" in text and "22, 80, 443, 8080, 8443" in text


def test_a_web_report_has_no_ports_line(reports_dir):
    run_to(None)
    assert store.report_data()["scope"]["allowed_ports"] == []
    assert "Allowed ports" not in render_html(store.report_data())
    if report.HAS_DOCX:
        store.generate_report("docx")
        assert "Allowed ports" not in _docx_text(reports_dir / "RCX-DEMO-001.docx")


def test_summary_counts_every_human_approval():
    run_to(None)
    summary = store.assessment_summary()
    assert summary.status == "Completed"
    assert summary.approvals == ["scope", "MEDIUM", "HIGH"]
    assert (summary.blocked, summary.confirmed, summary.for_review) == (1, 1, 2)
