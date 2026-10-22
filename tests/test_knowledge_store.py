"""The knowledge store: each file type, every refusal, the limits, the modes and redaction."""

import io
import json
import os
import stat

import pytest

from reconix import store
from reconix.store import knowledge
from reconix.store.knowledge import files, paths

from .support import TEST_PASSWORD


# --- fixtures for the binary file types ----------------------------------------------------
def _make_docx(path, paragraphs, table=None):
    from docx import Document

    document = Document()
    for para in paragraphs:
        document.add_paragraph(para)
    if table:
        t = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    document.save(str(path))


def _pdf_bytes(text):
    """A minimal one-page PDF with a single text line (pypdf extracts it)."""
    content = b"BT /F1 24 Tf 72 700 Td (" + text.encode("latin-1") + b") Tj ET\n"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"endstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(str(i).encode() + b" 0 obj\n" + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 " + str(len(objects) + 1).encode() + b"\n0000000000 65535 f \n")
    for off in offsets:
        out.write(("%010d 00000 n \n" % off).encode())
    out.write(b"trailer\n<< /Size " + str(len(objects) + 1).encode() + b" /Root 1 0 R >>\n")
    out.write(b"startxref\n" + str(xref).encode() + b"\n%%EOF")
    return out.getvalue()


# --- the file types ------------------------------------------------------------------------
def test_add_txt(tmp_path):
    p = tmp_path / "notes.txt"
    p.write_text("SQL injection was found in the search endpoint, marked high severity.")
    source = store.add_source(str(p))
    assert source.kind == "txt" and source.chunks >= 1 and source.bytes == p.stat().st_size
    assert [s.name for s in store.list_sources()] == ["notes.txt"]


def test_add_md(tmp_path):
    p = tmp_path / "runbook.md"
    p.write_text("# Runbook\n\nRotate the API keys after every penetration test engagement.")
    assert store.add_source(str(p)).kind == "md"


def test_add_json_stores_pretty_text(tmp_path):
    p = tmp_path / "config.json"
    p.write_text(json.dumps({"endpoint": "/search", "severity": "high"}))
    source = store.add_source(str(p))
    assert source.kind == "json"
    chunks = files.read_chunks(source.id)
    assert "endpoint" in chunks[0] and "/search" in chunks[0]


def test_add_yaml(tmp_path):
    p = tmp_path / "scope.yaml"
    p.write_text("target: staging.example.com\nmethods:\n  - GET\n  - POST\n")
    source = store.add_source(str(p))
    assert source.kind == "yaml"
    assert "staging.example.com" in files.read_chunks(source.id)[0]


def test_add_docx_paragraphs_and_tables(tmp_path):
    p = tmp_path / "report.docx"
    _make_docx(p, ["Executive summary of the engagement."],
               table=[["Severity", "Count"], ["High", "3"]])
    source = store.add_source(str(p))
    assert source.kind == "docx"
    text = " ".join(files.read_chunks(source.id))
    assert "Executive summary" in text and "Severity | Count" in text and "High | 3" in text


def test_add_pdf_page_text(tmp_path):
    p = tmp_path / "paper.pdf"
    p.write_bytes(_pdf_bytes("SQL injection reference material"))
    source = store.add_source(str(p))
    assert source.kind == "pdf"
    assert "SQL injection reference material" in files.read_chunks(source.id)[0]


# --- the refusals --------------------------------------------------------------------------
def test_refuse_empty_path():
    with pytest.raises(store.StoreValidationError):
        store.add_source("   ")


def test_refuse_missing_file(tmp_path):
    with pytest.raises(store.StoreValidationError):
        store.add_source(str(tmp_path / "nope.txt"))


def test_refuse_directory(tmp_path):
    with pytest.raises(store.StoreValidationError, match="regular file"):
        store.add_source(str(tmp_path))


def test_refuse_symlink(tmp_path):
    real = tmp_path / "real.txt"
    real.write_text("secret runbook contents here")
    link = tmp_path / "link.txt"
    link.symlink_to(real)
    with pytest.raises(store.StoreValidationError, match="symlink"):
        store.add_source(str(link))


def test_refuse_unknown_extension(tmp_path):
    p = tmp_path / "tool.exe"
    p.write_bytes(b"MZ binary")
    with pytest.raises(store.StoreValidationError, match="Unsupported"):
        store.add_source(str(p))


def test_refuse_too_large(tmp_path):
    p = tmp_path / "huge.txt"
    with open(p, "wb") as handle:
        handle.seek(10 * 1024 * 1024 + 1)
        handle.write(b"x")
    with pytest.raises(store.StoreValidationError, match="too large"):
        store.add_source(str(p))


def test_refuse_invalid_json(tmp_path):
    p = tmp_path / "bad.json"
    p.write_text("{not valid json,,,}")
    with pytest.raises(store.StoreValidationError, match="valid JSON"):
        store.add_source(str(p))


def test_refuse_encrypted_pdf(tmp_path):
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter()
    writer.append(PdfReader(io.BytesIO(_pdf_bytes("secret page"))))
    writer.encrypt("open-sesame")
    p = tmp_path / "locked.pdf"
    with open(p, "wb") as handle:
        writer.write(handle)
    with pytest.raises(store.StoreValidationError, match="encrypted"):
        store.add_source(str(p))


def test_refuse_unreadable_pdf(tmp_path):
    p = tmp_path / "broken.pdf"
    p.write_bytes(b"%PDF-1.4 this is not really a pdf")
    with pytest.raises(store.StoreValidationError, match="PDF"):
        store.add_source(str(p))


def test_refuse_file_with_no_text(tmp_path):
    p = tmp_path / "empty.txt"
    p.write_text("   \n\n   ")
    with pytest.raises(store.StoreValidationError, match="no readable text"):
        store.add_source(str(p))


# --- the limits ----------------------------------------------------------------------------
def test_source_limit(tmp_path):
    for i in range(files.MAX_SOURCES):
        p = tmp_path / f"f{i}.txt"
        p.write_text(f"reference document number {i} about scope and findings")
        store.add_source(str(p))
    assert len(store.list_sources()) == files.MAX_SOURCES
    over = tmp_path / "over.txt"
    over.write_text("one more document")
    with pytest.raises(store.StoreValidationError, match="limit"):
        store.add_source(str(over))


def test_passage_cap(tmp_path, monkeypatch):
    # A document with more paragraphs than the cap keeps at most MAX_PASSAGES.
    import sys

    import reconix.store.knowledge.chunk  # noqa: F401  (ensure the submodule is imported)
    chunk_mod = sys.modules["reconix.store.knowledge.chunk"]
    monkeypatch.setattr(chunk_mod, "MAX_PASSAGES", 5)
    big = "\n\n".join(f"paragraph {i} " + "word " * 300 for i in range(20))
    p = tmp_path / "big.txt"
    p.write_text(big)
    source = store.add_source(str(p))
    assert source.chunks == 5


# --- the modes -----------------------------------------------------------------------------
def test_folder_and_files_are_private(tmp_path):
    p = tmp_path / "notes.txt"
    p.write_text("reference material for the assessment")
    source = store.add_source(str(p))
    folder = paths.knowledge_dir()
    assert stat.S_IMODE(os.stat(folder).st_mode) == 0o700
    assert stat.S_IMODE(os.stat(files.index_path()).st_mode) == 0o600
    assert stat.S_IMODE(os.stat(files.source_path(source.id)).st_mode) == 0o600


# --- redaction at ingest -------------------------------------------------------------------
def test_secret_is_masked_on_disk(tmp_path):
    p = tmp_path / "creds.txt"
    p.write_text(f"The admin login uses password={TEST_PASSWORD} on the staging box.")
    source = store.add_source(str(p))
    on_disk = files.source_path(source.id).read_text(encoding="utf-8")
    assert TEST_PASSWORD not in on_disk
    assert TEST_PASSWORD not in files.index_path().read_text(encoding="utf-8")
    assert TEST_PASSWORD not in " ".join(files.read_chunks(source.id))


# --- remove / clear ------------------------------------------------------------------------
def test_remove_source(tmp_path):
    p = tmp_path / "a.txt"
    p.write_text("one reference document")
    source = store.add_source(str(p))
    store.remove_source(source.id)
    assert store.list_sources() == []
    assert not files.source_path(source.id).exists()


def test_remove_unknown_raises(tmp_path):
    with pytest.raises(store.StoreValidationError, match="Unknown add-on"):
        store.remove_source("0" * 32)


def test_clear_removes_everything(tmp_path):
    for name in ("a.txt", "b.txt"):
        p = tmp_path / name
        p.write_text(f"document {name}")
        store.add_source(str(p))
    store.clear_knowledge()
    assert store.list_sources() == []
    assert not files.index_path().exists()


# --- a corrupt index -----------------------------------------------------------------------
def test_corrupt_index_gives_a_clear_error(tmp_path):
    paths.ensure_dir()
    files.index_path().write_text("this is not json at all", encoding="utf-8")
    with pytest.raises(store.StoreValidationError, match="corrupt"):
        store.list_sources()
    # search never crashes the chat on a corrupt index.
    assert store.search_knowledge("anything") == []


def test_wrong_format_string_is_corrupt(tmp_path):
    paths.ensure_dir()
    files.index_path().write_text(json.dumps({"format": "other/v9", "sources": []}),
                                  encoding="utf-8")
    with pytest.raises(store.StoreValidationError, match="corrupt"):
        store.list_sources()


# --- search across sources -----------------------------------------------------------------
def test_search_returns_passage_objects(tmp_path):
    corpus = {
        "runbook.md": "SQL injection was found in the search endpoint and marked high severity.",
        "policy.txt": "Scope approval requires a signed authorization from the system owner.",
        "notes.md": "The login form accepts an email and a password over HTTPS only.",
    }
    for name, text in corpus.items():
        p = tmp_path / name
        p.write_text(text)
        store.add_source(str(p))
    hits = store.search_knowledge("sql injection severity endpoint search")
    assert hits and hits[0].source == "runbook.md"
    assert isinstance(hits[0], knowledge.Passage)
    assert hits[0].score > 0
